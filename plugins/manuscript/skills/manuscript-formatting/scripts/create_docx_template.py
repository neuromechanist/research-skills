#!/usr/bin/env python3
"""Create a clean, editable Pandoc reference DOCX for manuscript drafts."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(89, 89, 89)


def set_font(style, *, name: str, size: float, bold=None, italic=None, color=BLACK):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{slot}"), name)


def ensure_paragraph_style(styles, name: str, base: str = "Normal"):
    try:
        return styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles[base]
        return style


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, value, end))
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def add_sample_table(doc):
    caption = doc.add_paragraph(style="TableCaption")
    lead = caption.add_run("Table 1. ")
    lead.bold = True
    caption.add_run("A concise title that explains what the table summarizes.")

    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = (Inches(1.35), Inches(2.8), Inches(2.2))
    for column, width in zip(table.columns, widths):
        column.width = width
    rows = (
        ("Stage", "Action", "Evidence to retain"),
        ("Before", "Record a shared event", "Offset and jitter"),
        ("After", "Inspect aligned streams", "Timing summary"),
    )
    for row_index, values in enumerate(rows):
        for column_index, (cell, value) in enumerate(zip(table.rows[row_index].cells, values)):
            cell.width = widths[column_index]
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.style = "Compact"
                for run in paragraph.runs:
                    run.font.name = "Arial" if row_index == 0 else "Georgia"
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "4A4A4A")
                cell._tc.get_or_add_tcPr().append(shading)
            elif row_index == 2:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F2F2F2")
                cell._tc.get_or_add_tcPr().append(shading)


def create_template(output: Path, title_text: str):
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)

    styles = doc.styles
    normal = styles["Normal"]
    set_font(normal, name="Georgia", size=11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.14
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.first_line_indent = Inches(0)

    title = styles["Title"]
    set_font(title, name="Arial", size=23, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(9)
    title.paragraph_format.line_spacing = 1.02
    title.paragraph_format.keep_with_next = True
    title_ppr = title.element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)

    subtitle = styles["Subtitle"]
    set_font(subtitle, name="Arial", size=10.5, color=MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(8)
    subtitle.paragraph_format.keep_with_next = True

    for name, size, before, after in (
        ("Heading 1", 15, 16, 6),
        ("Heading 2", 12.5, 12, 4),
        ("Heading 3", 11.5, 9, 3),
    ):
        style = styles[name]
        set_font(style, name="Arial", size=size, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.05
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.widow_control = True

    for name, size, before, after, align in (
        ("Author", 11, 8, 2, WD_ALIGN_PARAGRAPH.LEFT),
        ("Affiliation", 10, 0, 2, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        style = ensure_paragraph_style(styles, name)
        set_font(style, name="Arial", size=size, color=MUTED if name == "Affiliation" else BLACK)
        style.paragraph_format.alignment = align
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    figure_image = ensure_paragraph_style(styles, "FigureImage")
    set_font(figure_image, name="Georgia", size=11)
    figure_image.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_image.paragraph_format.space_before = Pt(8)
    figure_image.paragraph_format.space_after = Pt(0)
    figure_image.paragraph_format.keep_with_next = True
    figure_image.paragraph_format.keep_together = True

    image_caption = ensure_paragraph_style(styles, "ImageCaption")
    set_font(image_caption, name="Georgia", size=9.5)
    image_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    image_caption.paragraph_format.space_before = Pt(4)
    image_caption.paragraph_format.space_after = Pt(13)
    image_caption.paragraph_format.line_spacing = 1.08
    image_caption.paragraph_format.keep_together = True
    image_caption.paragraph_format.widow_control = True

    table_caption = ensure_paragraph_style(styles, "TableCaption")
    set_font(table_caption, name="Georgia", size=9.5)
    table_caption.paragraph_format.space_before = Pt(10)
    table_caption.paragraph_format.space_after = Pt(4)
    table_caption.paragraph_format.line_spacing = 1.08
    table_caption.paragraph_format.keep_with_next = True
    table_caption.paragraph_format.keep_together = True

    bibliography = ensure_paragraph_style(styles, "Bibliography")
    set_font(bibliography, name="Georgia", size=9.5)
    bibliography.paragraph_format.left_indent = Inches(0.28)
    bibliography.paragraph_format.first_line_indent = Inches(-0.28)
    bibliography.paragraph_format.space_before = Pt(0)
    bibliography.paragraph_format.space_after = Pt(2)
    bibliography.paragraph_format.line_spacing = 1.0
    bibliography.paragraph_format.widow_control = True

    compact = ensure_paragraph_style(styles, "Compact", base="Body Text")
    set_font(compact, name="Georgia", size=9)
    compact.paragraph_format.space_before = Pt(0)
    compact.paragraph_format.space_after = Pt(2)
    compact.paragraph_format.line_spacing = 1.05
    compact.paragraph_format.widow_control = True

    for name in ("Body Text", "Block Text"):
        try:
            style = styles[name]
        except KeyError:
            continue
        set_font(style, name="Georgia", size=11)
        style.paragraph_format.line_spacing = 1.14
        style.paragraph_format.space_after = Pt(7)

    # Pandoc uses the Table style identifier; define compact cell spacing.
    try:
        table_style = styles["Table"]
    except KeyError:
        table_style = styles.add_style("Table", WD_STYLE_TYPE.TABLE)
    try:
        table_style.base_style = styles["Table Normal"]
    except KeyError:
        pass
    table_properties = table_style.element.find(qn("w:tblPr"))
    if table_properties is None:
        table_properties = OxmlElement("w:tblPr")
        table_style.element.append(table_properties)
    margins = table_properties.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        table_properties.append(margins)
    for side, value in (("top", 55), ("bottom", 55), ("left", 90), ("right", 90)):
        margin = margins.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    add_page_field(footer)

    paragraph = doc.add_paragraph(style="Title")
    paragraph.add_run(title_text)
    paragraph = doc.add_paragraph(style="Subtitle")
    paragraph.add_run("Working manuscript draft")
    doc.add_paragraph("Author Name(s)", style="Author")
    doc.add_paragraph("Department, Institution | ORCID", style="Affiliation")
    doc.add_paragraph("Corresponding author: name@example.org", style="Affiliation")

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "Summarize the question, approach, key findings, and implication within the journal's word limit."
    )
    keyword = doc.add_paragraph()
    run = keyword.add_run("Keywords: ")
    run.bold = True
    keyword.add_run("keyword one; keyword two; keyword three")

    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(
        "State the problem, the gap, and the study objective or hypothesis."
    )
    doc.add_heading("Methods", level=1)
    doc.add_paragraph(
        "Describe the design, participants or materials, procedures, measures, and analysis in enough detail to reproduce the work."
    )
    doc.add_heading("Results", level=1)
    doc.add_paragraph(
        "Present findings in research-question order, citing figures and tables. Reserve interpretation for the Discussion."
    )
    figure = doc.add_table(rows=1, cols=1)
    figure.style = "Table Grid"
    figure.autofit = False
    figure.columns[0].width = Inches(6.35)
    cell = figure.cell(0, 0)
    cell.width = Inches(6.35)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "F2F2F2")
    cell._tc.get_or_add_tcPr().append(shade)
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), "950")
    figure.rows[0]._tr.get_or_add_trPr().append(height)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Figure placeholder: replace with an image in the Markdown source")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    caption = doc.add_paragraph(style="ImageCaption")
    run = caption.add_run("Figure 1. ")
    run.bold = True
    caption.add_run("State the main message; define symbols, abbreviations, and visual encodings.")

    add_sample_table(doc)
    doc.add_heading("Discussion", level=1)
    doc.add_paragraph(
        "Interpret findings, limitations, and implications; distinguish evidence-based conclusions from future directions."
    )
    doc.add_heading("References", level=1)
    doc.add_paragraph(
        "Replace with references rendered by Pandoc citeproc or formatted for the target journal.",
        style="Bibliography",
    )
    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("Manuscript_Draft_Reference.docx"),
        help="where to write the reference DOCX (default: ./Manuscript_Draft_Reference.docx)",
    )
    parser.add_argument(
        "--title",
        default="Manuscript Title Goes Here",
        help="sample title shown when the reference file is opened in Word",
    )
    args = parser.parse_args()
    create_template(args.output, args.title)
    print(f"Wrote reference DOCX: {args.output}")


if __name__ == "__main__":
    main()
